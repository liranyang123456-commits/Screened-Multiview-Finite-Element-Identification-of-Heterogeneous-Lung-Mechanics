from pathlib import Path
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor

src_dir = Path(sys.argv[1])
out_dir = Path(sys.argv[2])
out_dir.mkdir(parents=True, exist_ok=True)

LABELS = ('**Author Response:**', '**Location of revisions:**')
QUOTE_LABELS = ("Reviewer's Comment:", "Editor's Comment:")

md_paths = sorted(src_dir.glob('Response_Reviewer[1-4].md')) + sorted(src_dir.glob('Response_to_Editor.md'))
for md_path in md_paths:
    text = md_path.read_text(encoding='utf-8-sig')
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for stylename, size in [('Title', 16), ('Heading 1', 13), ('Heading 2', 11.5)]:
        style = doc.styles[stylename]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 78, 121)
    if '|' in text[:80]:
        text = re.sub(r'^\s*\|+', '', text)
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        heading = re.match(r'^(#{1,3}) ', stripped)
        if heading:
            level = len(heading.group(1)) - 1
            h = doc.add_heading(stripped[heading.end():].strip(), level=level)
            h.paragraph_format.keep_with_next = True
            continue
        if stripped.startswith('> '):
            body = stripped[2:].replace('  ', ' ')
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.18)
            for part in re.split(r'(\*\*.*?\*\*)', body):
                if not part:
                    continue
                bold = len(part) > 4 and part.startswith('**') and part.endswith('**')
                r = p.add_run(part[2:-2] if bold else part)
                r.italic = True
                if bold:
                    r.bold = True
                r.font.color.rgb = RGBColor(90, 90, 90)
            if body.replace('**', '').strip() in QUOTE_LABELS:
                p.paragraph_format.keep_with_next = True
            continue
        if stripped == '---':
            continue
        p = doc.add_paragraph()
        if stripped in LABELS:
            p.paragraph_format.keep_with_next = True
        parts = re.split(r'(\*\*.*?\*\*|`[^`]+`)', stripped)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                p.add_run(part[2:-2]).bold = True
            elif part.startswith('`') and part.endswith('`'):
                run = p.add_run(part[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
            else:
                p.add_run(part)
    output = out_dir / f'{md_path.stem}.docx'
    props = doc.core_properties
    props.author = props.last_modified_by = 'Ranyang Li'
    props.comments = ''
    props.title = md_path.stem.replace('_', ' ')
    doc.save(output)
    print(output)
